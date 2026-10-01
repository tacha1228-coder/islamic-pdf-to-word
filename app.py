import streamlit as st
import io
import tempfile
import os
import re
import time

import pdfplumber
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

st.set_page_config(page_title="محول PDF إلى Word", page_icon="📖", layout="wide")

st.error("""
### ⚠️ تنبيه شرعي وأخلاقي هام

**لا أُحل استخدام هذا التطبيق في:**
- ❌ السرقات العلمية أو انتحال المحتوى
- ❌ تفريغ محتويات الكتب التي يُخاف فيها شرع الله
- ❌ أي استخدام يخالف الأمانة العلمية

**من يفعل ذلك فإني خصيمه يوم القيامة، ولا سماح بيننا في الدنيا ولا في الآخرة.**
""")

st.title("📖 محول PDF إلى Word")
st.markdown("**مخصص لتفريغ دروس العلماء والمحتوى الديني** بأمانة علمية كاملة")

with st.sidebar:
    st.header("⚙️ الإعدادات")
    font_choice = st.selectbox("الخط العربي", ["Traditional Arabic", "Amiri", "Arial", "Scheherazade New"])
    font_size = st.slider("حجم الخط", 10, 20, 14)
    add_page_numbers = st.checkbox("إضافة أرقام الصفحات", value=True)
    add_source = st.checkbox("إضافة اسم الملف الأصلي", value=True)

uploaded = st.file_uploader("اختاري ملف PDF", type=["pdf"])

def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def clean_text(text):
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def format_time(seconds):
    if seconds < 60:
        return f"{int(seconds)} ثانية"
    elif seconds < 3600:
        m = int(seconds // 60)
        s = int(seconds % 60)
        return f"{m} د {s} ث"
    else:
        h = int(seconds // 3600)
        m = int((seconds % 3600) // 60)
        return f"{h} س {m} د"

if uploaded:
    file_size_mb = uploaded.size / (1024 * 1024)
    st.success(f"✅ تم رفع الملف: {uploaded.name} ({file_size_mb:.2f} MB)")
    
    if st.button("🚀 ابدأ التحويل", type="primary"):
        try:
            # ================= المرحلة 1: حفظ الملف =================
            st.markdown("### 📤 المرحلة 1 من 3: تجهيز الملف")
            bar1 = st.progress(0, text="0%")
            status1 = st.empty()
            
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
            tmp.write(uploaded.read())
            tmp.close()
            
            for i in range(1, 11):
                bar1.progress(i * 10, text=f"{i*10}% - حفظ الملف في السيرفر")
                time.sleep(0.05)
            status1.success("✅ تم تجهيز الملف بنجاح")
            
            # ================= المرحلة 2: استخراج النص =================
            st.markdown("### ⚙️ المرحلة 2 من 3: استخراج النص من الصفحات")
            bar2 = st.progress(0, text="0%")
            status2 = st.empty()
            
            pages_data = []
            pdf_metadata = {}
            
            with pdfplumber.open(tmp.name) as pdf:
                pdf_metadata = pdf.metadata or {}
                total = len(pdf.pages)
                start_time = time.time()
                
                for i, page in enumerate(pdf.pages):
                    text = page.extract_text(x_tolerance=2, y_tolerance=2) or ""
                    pages_data.append(text)
                    
                    percent = int(((i + 1) / total) * 100)
                    elapsed = time.time() - start_time
                    if i > 0:
                        avg = elapsed / (i + 1)
                        remaining = avg * (total - i - 1)
                        time_str = f" | ⏳ الوقت المتبقي: {format_time(remaining)}"
                    else:
                        time_str = " | ⏳ جاري الحساب..."
                    
                    bar2.progress((i + 1) / total, text=f"{percent}% - الصفحة {i+1} من {total}{time_str}")
            
            all_text = "\n\n".join([p for p in pages_data if p.strip()])
            
            if not all_text.strip():
                st.warning("⚠️ لا يوجد نص قابل للاستخراج. قد يكون الملف صورًا ممسوحة ضوئيًا.")
                st.stop()
            
            status2.success(f"✅ تم استخراج {total} صفحة بنجاح")
            
            # ================= المرحلة 3: بناء ملف Word =================
            st.markdown("### 📥 المرحلة 3 من 3: إنشاء ملف Word")
            bar3 = st.progress(0, text="0%")
            status3 = st.empty()
            
            bar3.progress(10, text="10% - تجهيز المستند")
            time.sleep(0.1)
            
            doc = Document()
            style = doc.styles['Normal']
            style.font.name = font_choice
            style.font.size = Pt(font_size)
            style.element.rPr.rFonts.set(qn('w:eastAsia'), font_choice)
            style.element.rPr.rFonts.set(qn('w:cs'), font_choice)
            
            bar3.progress(20, text="20% - إضافة معلومات المصدر")
            
            if add_source:
                title = doc.add_heading(f"المصدر: {uploaded.name}", level=1)
                set_rtl(title)
                if pdf_metadata.get('Title'):
                    p = doc.add_paragraph(f"العنوان الأصلي: {pdf_metadata['Title']}")
                    set_rtl(p)
                if pdf_metadata.get('Author'):
                    p = doc.add_paragraph(f"المؤلف: {pdf_metadata['Author']}")
                    set_rtl(p)
                doc.add_paragraph("—" * 30)
            
            for i, page_text in enumerate(pages_data, start=1):
                if not page_text.strip():
                    continue
                
                if add_page_numbers:
                    page_marker = doc.add_paragraph()
                    set_rtl(page_marker)
                    run = page_marker.add_run(f"【 صفحة {i} 】")
                    run.bold = True
                    run.font.color.rgb = RGBColor(0x88, 0x44, 0x00)
                
                cleaned = clean_text(page_text)
                for line in cleaned.split('\n'):
                    if line.strip():
                        p = doc.add_paragraph(line)
                        set_rtl(p)
                doc.add_paragraph()
                
                percent = 20 + int((i / total) * 70)
                bar3.progress(percent / 100, text=f"{percent}% - كتابة الصفحة {i} من {total}")
            
            bar3.progress(95, text="95% - حفظ الملف النهائي")
            buf = io.BytesIO()
            doc.save(buf)
            buf.seek(0)
            bar3.progress(100, text="100% - ✅ الملف جاهز!")
            status3.success("✅ تم إنشاء ملف Word بنجاح")
            
            os.unlink(tmp.name)
            
            # ================= التحميل =================
            st.divider()
            st.success("🎉 اكتمل التحويل! يمكنك تحميل الملف الآن:")
            st.download_button(
                "📄 تحميل ملف Word",
                buf,
                f"converted_{uploaded.name.replace('.pdf', '')}.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            st.info("💡 الملف نصي بالكامل — انسخي مع ذكر المصدر للأمانة العلمية")
            
        except Exception as e:
            st.error(f"حدث خطأ: {e}")
